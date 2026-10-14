import pytest

from copilot.ingest.convert import import_document
from copilot.ingest.loaders import load_markdown
from copilot.models import Role


def _docx(path, title="Travel Policy 2027"):
    import docx

    d = docx.Document()
    d.core_properties.title = title
    d.add_heading("Booking", level=1)
    d.add_paragraph("Book flights at least 14 days ahead.")
    d.add_paragraph("Use the travel portal", style="List Bullet")
    d.add_paragraph("Keep the receipt", style="List Bullet")
    d.add_heading("Hotels", level=2)
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "City", "Nightly cap"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "London", "$250"
    d.save(str(path))
    return path


def _pdf(path, pages):
    """Write a minimal PDF with one text line per entry, one page per list."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None]
    kids = []
    font_id = 3
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for lines in pages:
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td"]
        for line in lines:
            ops += [f"({line}) Tj", "T*"]
        ops.append("ET")
        stream = "\n".join(ops)
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objs)
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R"
            f" /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        )
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))
    return path


def test_docx_becomes_tagged_markdown(tmp_path):
    src = _docx(tmp_path / "Travel Policy.docx")
    out = import_document(src, tmp_path / "corpus", "policy", "manager")

    assert out == tmp_path / "corpus" / "policy" / "travel-policy.md"
    doc = load_markdown(out, "policy")
    assert doc.doc_id == "travel-policy"
    assert doc.title == "Travel Policy 2027"
    assert doc.min_role == Role.MANAGER
    assert doc.metadata["source"] == "Travel Policy.docx"
    headings = [s.heading for s in doc.sections]
    assert "Travel Policy 2027 > Booking" in headings
    assert "Travel Policy 2027 > Booking > Hotels" in headings
    text = out.read_text()
    assert "- Use the travel portal\n- Keep the receipt" in text
    assert "| London | $250 |" in text


def test_pdf_numbered_headings_become_sections(tmp_path):
    src = _pdf(
        tmp_path / "leave.pdf",
        [["1. Annual leave", "Employees get 25 days a year.", "2. Sick leave", "Up to 10 days."]],
    )
    out = import_document(src, tmp_path / "corpus", "policy", "employee", title="Leave Policy")
    headings = [s.heading for s in load_markdown(out, "policy").sections]
    assert headings == ["Leave Policy > 1. Annual leave", "Leave Policy > 2. Sick leave"]


def test_pdf_without_headings_is_split_by_page(tmp_path):
    src = _pdf(tmp_path / "notes.pdf", [["first page text"], ["second page text"]])
    out = import_document(src, tmp_path / "corpus", "engineering", "employee")
    headings = [s.heading for s in load_markdown(out, "engineering").sections]
    assert headings == ["Notes > Page 1", "Notes > Page 2"]


def test_refuses_to_overwrite_without_force(tmp_path):
    src = _docx(tmp_path / "policy.docx")
    corpus = tmp_path / "corpus"
    import_document(src, corpus, "policy", "employee")
    with pytest.raises(FileExistsError):
        import_document(src, corpus, "policy", "employee")
    import_document(src, corpus, "policy", "admin", overwrite=True)
    assert load_markdown(corpus / "policy" / "policy.md", "policy").min_role == Role.ADMIN


def test_refuses_a_doc_id_used_by_another_file(tmp_path):
    corpus = tmp_path / "corpus"
    (corpus / "policy").mkdir(parents=True)
    (corpus / "policy" / "old.md").write_text("---\ndoc_id: travel\nmin_role: employee\n---\nx")
    with pytest.raises(FileExistsError, match="already used"):
        import_document(_docx(tmp_path / "travel.docx"), corpus, "policy", "employee")


def test_rejects_bad_inputs(tmp_path):
    src = _docx(tmp_path / "a.docx")
    with pytest.raises(ValueError, match="unknown role"):
        import_document(src, tmp_path, "policy", "boss")
    with pytest.raises(ValueError, match="tenant"):
        import_document(src, tmp_path, "../etc", "employee")
    (tmp_path / "old.doc").write_bytes(b"x")
    with pytest.raises(ValueError, match="docx"):
        import_document(tmp_path / "old.doc", tmp_path, "policy", "employee")
