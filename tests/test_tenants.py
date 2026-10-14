"""Creating tenants and giving users access to them."""

import pytest
from fastapi.testclient import TestClient

from copilot.api import create_app
from copilot.cli import main
from copilot.ingest.convert import import_document
from copilot.ingest.loaders import discover
from copilot.tenants import create_tenant, list_tenants
from copilot.users import UserStore

PASSWORD = "LongEnough123"


@pytest.fixture
def corpus(tmp_path):
    root = tmp_path / "corpus"
    (root / "policy").mkdir(parents=True)
    (root / "policy" / "a.md").write_text("---\nmin_role: employee\n---\n# A\n\ntext")
    return root


@pytest.fixture
def users(tmp_path):
    store = UserStore(tmp_path / "users.db", iterations=1_000)
    store.create("alice", PASSWORD, "employee", ["policy"], must_change=False)
    store.create("bob", PASSWORD, "manager", ["policy"], must_change=False)
    return store


def test_create_and_list(corpus):
    tenant = create_tenant(corpus, "legal", "Legal team")
    assert tenant.label == "Legal team"
    listed = {t.name: (t.label, t.documents) for t in list_tenants(corpus)}
    assert listed == {"legal": ("Legal team", 0), "policy": ("Company policy", 1)}


def test_default_label_from_name(corpus):
    assert create_tenant(corpus, "hr-ops").label == "Hr Ops"


@pytest.mark.parametrize("name", ["Legal", "../x", "", "a b", "-x"])
def test_rejects_bad_names(corpus, name):
    with pytest.raises(ValueError):
        create_tenant(corpus, name)


def test_rejects_existing(corpus):
    with pytest.raises(FileExistsError):
        create_tenant(corpus, "policy")


def test_tenant_file_is_not_indexed(corpus, tmp_path):
    create_tenant(corpus, "legal")
    src = tmp_path / "x.docx"
    import docx

    d = docx.Document()
    d.add_heading("Contracts", level=1)
    d.add_paragraph("All contracts need two signatures.")
    d.save(str(src))
    import_document(src, corpus, "legal", "employee")
    docs = {(d.tenant, d.doc_id) for d in discover(corpus)}
    assert docs == {("policy", "a"), ("legal", "x")}


def test_set_tenants(users):
    assert users.set_tenants("alice", ["policy", "legal", "policy"]).tenants == ["policy", "legal"]
    with pytest.raises(ValueError):
        users.set_tenants("alice", [])
    with pytest.raises(KeyError):
        users.set_tenants("nobody", ["policy"])


def _cli(monkeypatch, users, *argv):
    monkeypatch.setenv("COPILOT_USERS_DB", str(users.path))
    monkeypatch.setenv("COPILOT_PASSWORD_HASH_ITERATIONS", "1000")
    from copilot.config import get_settings

    get_settings.cache_clear()
    try:
        main(list(argv))
    finally:
        get_settings.cache_clear()


def test_cli_add_tenant_with_grant(monkeypatch, capsys, corpus, users):
    _cli(monkeypatch, users, "tenants", "add", "legal", "--label", "Legal",
         "--grant", "alice,nobody", "--corpus-dir", str(corpus))  # fmt: skip
    out = capsys.readouterr().out
    assert "gave alice access" in out and "no such user or role: nobody" in out
    assert users.get("alice").tenants == ["policy", "legal"]
    assert users.get("bob").tenants == ["policy"]


def test_cli_grant_all_and_revoke(monkeypatch, corpus, users):
    _cli(monkeypatch, users, "tenants", "add", "legal", "--grant-all", "--corpus-dir", str(corpus))
    assert users.get("bob").tenants == ["policy", "legal"]
    _cli(monkeypatch, users, "users", "revoke", "bob", "policy")
    assert users.get("bob").tenants == ["legal"]
    with pytest.raises(SystemExit):
        _cli(monkeypatch, users, "users", "revoke", "bob", "legal")


def test_cli_grant_checks_tenant_exists(monkeypatch, corpus, users):
    with pytest.raises(SystemExit, match="no such tenant"):
        _cli(monkeypatch, users, "users", "grant", "alice", "legal", "--corpus-dir", str(corpus))
    create_tenant(corpus, "legal")
    _cli(monkeypatch, users, "users", "grant", "alice", "legal", "--corpus-dir", str(corpus))
    assert users.get("alice").tenants == ["policy", "legal"]


def test_tenants_endpoint_and_login(settings, service, corpus, users):
    create_tenant(corpus, "legal", "Legal")
    users.set_tenants("alice", ["policy", "legal"])
    client = TestClient(
        create_app(settings.model_copy(update={"corpus_dir": corpus}), service, users)
    )
    assert client.get("/tenants").json() == [
        {"name": "policy", "label": "Company policy"},
        {"name": "legal", "label": "Legal"},
    ]
    ok = client.post(
        "/auth/login", json={"username": "alice", "password": PASSWORD, "tenant": "legal"}
    )
    assert ok.status_code == 200 and ok.json()["user"]["tenant"] == "legal"
    denied = client.post(
        "/auth/login", json={"username": "bob", "password": PASSWORD, "tenant": "legal"}
    )
    assert denied.status_code == 403


def test_exist_ok_returns_existing(corpus):
    create_tenant(corpus, "legal", "Legal")
    assert create_tenant(corpus, "legal", "Other", exist_ok=True).label == "Legal"


def test_grant_by_role_covers_users_added_later(
    monkeypatch, capsys, settings, service, corpus, users
):
    _cli(monkeypatch, users, "tenants", "add", "legal", "--label", "Legal",
         "--grant", "admin,manager", "--corpus-dir", str(corpus))  # fmt: skip
    assert "open to every admin and manager" in capsys.readouterr().out
    (legal,) = [t for t in list_tenants(corpus) if t.name == "legal"]
    assert legal.roles == ["admin", "manager"]
    # Roles are kept in tenant.yaml, not copied onto accounts.
    assert users.get("bob").tenants == ["policy"]
    users.create("carl", PASSWORD, "manager", ["policy"], must_change=False)

    client = TestClient(
        create_app(settings.model_copy(update={"corpus_dir": corpus}), service, users)
    )

    def login(name):
        return client.post(
            "/auth/login", json={"username": name, "password": PASSWORD, "tenant": "legal"}
        )

    for manager in ("bob", "carl"):
        r = login(manager)
        assert r.status_code == 200
        assert r.json()["user"]["tenants"] == ["policy", "legal"]
    assert login("alice").status_code == 403  # employee

    # Taking the role off the tenant ends existing sessions in it.
    token = login("carl").json()["access_token"]
    (corpus / "legal" / "tenant.yaml").write_text("label: Legal\nroles: [admin]\n")
    r = client.get("/me", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 401


def test_grant_mixes_roles_and_users(monkeypatch, corpus, users):
    _cli(monkeypatch, users, "tenants", "add", "legal", "--grant", "admin,alice",
         "--corpus-dir", str(corpus))  # fmt: skip
    (legal,) = [t for t in list_tenants(corpus) if t.name == "legal"]
    assert legal.roles == ["admin"]
    assert users.get("alice").tenants == ["policy", "legal"]


def test_exist_ok_adds_roles(corpus):
    create_tenant(corpus, "legal", "Legal", ["admin"])
    t = create_tenant(corpus, "legal", roles=["manager", "admin"], exist_ok=True)
    assert t.roles == ["admin", "manager"]
    with pytest.raises(ValueError, match="unknown role"):
        create_tenant(corpus, "hr", roles=["boss"])
