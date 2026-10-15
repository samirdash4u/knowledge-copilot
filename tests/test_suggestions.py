"""Suggested questions under the chat box."""

from fastapi.testclient import TestClient

from copilot.api import create_app
from copilot.auth import issue_token
from copilot.models import Chunk, Principal, Role
from copilot.suggestions import configured, generated


def _get(settings, service, role, tenant):
    client = TestClient(create_app(settings, service))
    token = issue_token(settings, "u", role, tenant)
    return client.get("/suggestions", headers={"Authorization": f"Bearer {token}"})


def test_builtin_tenants_use_their_tenant_yaml(settings, service):
    r = _get(settings, service, "employee", "policy")
    assert r.status_code == 200
    assert r.json()[0] == "How many days of paid annual leave do I get?"
    assert "How do I roll back a service?" in _get(settings, service, "admin", "engineering").json()


def test_requires_sign_in(settings, service):
    assert TestClient(create_app(settings, service)).get("/suggestions").status_code == 401


def _chunk(doc, section, role, tenant="legal", title="Contracts Policy"):
    return Chunk(
        chunk_id=f"{tenant}/{doc}#{section}",
        doc_id=doc,
        tenant=tenant,
        title=title,
        section=f"{title} > {section}" if section else title,
        text="x",
        min_role=role,
    )


def test_generated_respects_role_and_tenant():
    chunks = [
        _chunk("contracts", "", Role.EMPLOYEE),
        _chunk("contracts", "Signatures", Role.EMPLOYEE),
        _chunk("contracts", "Settlement limits", Role.MANAGER),
        _chunk("nda", "Page 2", Role.EMPLOYEE, title="The NDA Template"),
        _chunk("other", "Secret", Role.EMPLOYEE, tenant="policy", title="Other"),
    ]
    employee = generated(Principal("e", Role.EMPLOYEE, "legal"), chunks)
    assert employee == [
        "What is the Contracts Policy about?",
        "What is The NDA Template about?",
        'What does the Contracts Policy say about "Signatures"?',
    ]
    manager = generated(Principal("m", Role.MANAGER, "legal"), chunks)
    # Material written for the caller's own role comes first.
    assert manager[0] == 'What does the Contracts Policy say about "Settlement limits"?'
    assert not any("Settlement" in q for q in employee)
    assert not any("Secret" in q for q in manager)


def test_generated_is_capped():
    chunks = [_chunk(f"d{i}", f"S{i}", Role.EMPLOYEE) for i in range(10)]
    assert len(generated(Principal("e", Role.EMPLOYEE, "legal"), chunks)) == 4


def test_configured_forms():
    assert configured(None, "employee") is None
    assert configured({"label": "x"}, "employee") is None
    assert configured(["A?", " ", "B?"], "admin") == ["A?", "B?"]
    per_role = {"manager": ["M?"]}
    assert configured(per_role, "manager") == ["M?"]
    assert configured(per_role, "employee") is None  # not set for this role: generated instead


# ---------------------------------------------------------------- add-doc writes them


def _doc(tmp_path):
    from copilot.ingest.loaders import load_markdown

    p = tmp_path / "brand.md"
    p.write_text(
        "---\nmin_role: employee\ntitle: Brand Guidelines\n---\n# Brand Guidelines\n\n"
        "## Logo usage\n\nKeep clear space around the logo at all times.\n\n"
        "## Social media\n\nPosts on company accounts need the marketing lead's approval.\n\n"
        "## Campaign budgets\n<!-- acl: manager -->\n\nCampaigns above $20,000 need CMO sign-off.\n"
    )
    return load_markdown(p, "marketing")


class FakeLLM:
    def __init__(self, replies):
        self.replies = list(replies)
        self.calls = 0

    def chat(self, messages):
        self.calls += 1
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        return reply


def test_document_questions_from_headings(tmp_path):
    from copilot.suggestions import document_questions

    by_role, used_llm = document_questions(_doc(tmp_path))
    assert not used_llm
    assert by_role["employee"] == [
        'What does the Brand Guidelines say about "Logo usage"?',
        'What does the Brand Guidelines say about "Social media"?',
    ]
    # The manager-only section is suggested to managers (first) and admins, never employees.
    assert by_role["manager"][0] == 'What does the Brand Guidelines say about "Campaign budgets"?'
    assert "Campaign budgets" in " ".join(by_role["admin"])


def test_document_questions_from_llm(tmp_path):
    from copilot.suggestions import document_questions

    llm = FakeLLM(
        [
            "Question: How much space should I leave around the logo?",
            "Here is a question you could ask about posts and approvals",  # not a question
            '"Who signs off on a large campaign?"',
        ]
    )
    by_role, used_llm = document_questions(_doc(tmp_path), llm)
    assert used_llm and llm.calls == 3
    assert by_role["employee"] == [
        "How much space should I leave around the logo?",
        'What does the Brand Guidelines say about "Social media"?',
    ]
    assert by_role["manager"][0] == "Who signs off on a large campaign?"


def test_unreachable_llm_falls_back_once(tmp_path):
    from copilot.suggestions import document_questions

    llm = FakeLLM([ConnectionError("no ollama")])
    by_role, used_llm = document_questions(_doc(tmp_path), llm)
    assert not used_llm and llm.calls == 1  # gave up after the first failure
    assert len(by_role["employee"]) == 2


def test_add_to_tenant_merges(tmp_path):
    import yaml

    from copilot.suggestions import add_to_tenant

    (tmp_path / "tenant.yaml").write_text(
        'label: Marketing\nroles: [manager]\nsuggestions: ["Old A?", "Old B?"]\n'
    )
    assert add_to_tenant(tmp_path, {"employee": ["New?", "Old A?"], "manager": ["M1?"]})
    meta = yaml.safe_load((tmp_path / "tenant.yaml").read_text())
    assert meta["label"] == "Marketing" and meta["roles"] == ["manager"]
    assert meta["suggestions"]["employee"] == ["New?", "Old A?", "Old B?"]
    assert meta["suggestions"]["manager"] == ["M1?", "Old A?", "Old B?"]
    assert meta["suggestions"]["admin"] == ["Old A?", "Old B?"]

    add_to_tenant(tmp_path, {"employee": [f"Q{i}?" for i in range(10)]})
    assert (
        len(yaml.safe_load((tmp_path / "tenant.yaml").read_text())["suggestions"]["employee"]) == 6
    )


def test_add_to_tenant_leaves_broken_yaml_alone(tmp_path):
    from copilot.suggestions import add_to_tenant

    (tmp_path / "tenant.yaml").write_text("suggestions: [What is X?]\n")
    assert not add_to_tenant(tmp_path, {"employee": ["New?"]})
    assert (tmp_path / "tenant.yaml").read_text() == "suggestions: [What is X?]\n"


def test_cli_add_doc_writes_suggestions(monkeypatch, capsys, tmp_path):
    import docx
    import yaml

    from copilot.cli import main
    from copilot.config import get_settings
    from copilot.tenants import create_tenant

    corpus = tmp_path / "corpus"
    create_tenant(corpus, "marketing", "Marketing")
    d = docx.Document()
    d.add_heading("Logo usage", level=1)
    d.add_paragraph("Keep clear space around the logo at all times.")
    d.save(str(tmp_path / "brand.docx"))

    monkeypatch.setenv("COPILOT_LLM", "echo")
    get_settings.cache_clear()
    try:
        main(["add-doc", str(tmp_path / "brand.docx"), "--tenant", "marketing",
              "--min-role", "employee", "--corpus-dir", str(corpus)])  # fmt: skip
        meta = yaml.safe_load((corpus / "marketing" / "tenant.yaml").read_text())
        assert meta["suggestions"]["employee"] == ['What does the Brand say about "Logo usage"?']
        assert "added suggested questions" in capsys.readouterr().out

        main(["add-doc", str(tmp_path / "brand.docx"), "--tenant", "marketing", "--force",
              "--min-role", "employee", "--corpus-dir", str(corpus), "--no-suggest"])  # fmt: skip
        assert "added suggested questions" not in capsys.readouterr().out
    finally:
        get_settings.cache_clear()
